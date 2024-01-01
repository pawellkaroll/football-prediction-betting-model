import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
import joblib
import time
import xgboost as xgb
import docx

from sklearn.linear_model import LogisticRegression
from sklearn.ensemble import RandomForestClassifier
from sklearn.preprocessing import MinMaxScaler, StandardScaler, LabelEncoder,Normalizer, PowerTransformer, MaxAbsScaler, LabelEncoder, OneHotEncoder
from sklearn.model_selection import GridSearchCV, cross_val_score, RepeatedStratifiedKFold
from sklearn.metrics import confusion_matrix, accuracy_score, recall_score, precision_score, f1_score, roc_curve, roc_auc_score, make_scorer,matthews_corrcoef, average_precision_score
from sklearn.feature_selection import mutual_info_classif,f_classif
from sklearn.feature_selection import VarianceThreshold, SelectKBest 
from sklearn.pipeline import Pipeline 
from pandas.plotting import scatter_matrix
from scipy.stats import boxcox
from numpy import mean, std
from collections import Counter
from matplotlib.colors import ListedColormap



# save dateframe to word file
def df_to_word(df,path):
    
    doc = docx.Document(path)

    t = doc.add_table(df.shape[0]+1, df.shape[1])

    for j in range(df.shape[-1]):
        t.cell(0,j).text = df.columns[j]

    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i+1,j).text = str(df.values[i,j])
    
    doc.add_page_break()
    doc.save(path)


# get train and test sets
def train_test_split(df,start,window,features='all',forecast=1):
    if features=='all':
        X=df.loc[:, df.columns != 'FTR']
    else:
        X=df.loc[:,df.columns.isin(features)]
    y=df[['FTR']]
    start=start-1
    X_train=X[start*10:start*10+window*10]
    X_test=X[start*10+window*10:(start*10+window*10)+10*forecast]
    y_train=y.loc[start*10:(start*10+window*10)-1,:]
    y_test=y.loc[start*10+window*10:(start*10+window*10)+10*forecast-1,:]
    return(X_train,X_test,y_train,y_test)


# get 1) list of probabilities of each option 2) list of predictions (outcome with highest probability)
def get_forecasts(df,features,stand_df=False, fct_num=8,fct_win=1, model='RandomForest'):

   if stand_df==True:
      X=df.drop(['HomeTeam',"AwayTeam","FTR"],axis=1)
      y=df[["FTR"]]
      tmp=df[['HomeTeam',"AwayTeam"]]
      scaler = StandardScaler()
      X_s = scaler.fit_transform(X)
      X_s= pd.DataFrame(X_s, columns = X.columns)
      df=pd.concat([tmp,y,X_s],axis=1)
   
   scores_list=[]
   cross_scores_list=[]
   prob_list=[]
   pred_list=[]
   
   for i in range(0,fct_num):
      X_train,X_test,y_train,y_test=train_test_split(df=df,start=1,window=30+i,forecast=fct_win,features=features) 
      
      if model=='RandomForest':
         model = RandomForestClassifier(random_state=44,class_weight='balanced')
      elif model=='LogisticRegression':
         model = LogisticRegression(random_state=44,multi_class='multinomial',solver='lbfgs',max_iter=10000,class_weight='balanced')
      elif model=='XGBoost':
         model = xgb.XGBClassifier(random_state=44)
      
      model.fit(X_train, y_train.values.ravel())

      predictions = model.predict(X_test)
      probabilities = model.predict_proba(X_test)
      pred_list.append(predictions)
      prob_list.append(probabilities)

      score = accuracy_score(y_test, predictions) 
      scores_list.append(score)
      cross_scores = cross_val_score(model, X_train,y_train.values.ravel(),scoring='accuracy',cv=10)
      avg_cross_score=mean(cross_scores)
      cross_scores_list.append(avg_cross_score)
      
      print('Model '+ str(i) +': '+'Accuracy of '+ str(score))
      print('Model '+ str(i) +': '+'Cross Validation Accuracy of '+ str(avg_cross_score))
   
   print('All models '+': '+'Average accuracy of '+ str(mean(scores_list)))
   print('All models '+': '+'Cross Validation Average accuracy of '+ str(mean(cross_scores_list)))
   return(prob_list, pred_list)


# get list of predictions 
def binomial_x3(df,features,model="RandomForest",variables=['AvgH','AvgA']):

    # Home / Not Home
    df_h=df.replace({'A': 'NH', 'D': 'NH'})
    df_h=df_h[variables]
    if model=='XGBoost':
        le = LabelEncoder()
        df_h['FTR'] = le.fit_transform(df_h['FTR'])
    probs,preds=get_forecasts(df=df_h,features=features,stand_df=False,fct_num=8,fct_win=1,model=model)

    df_prob_h=pd.DataFrame(np.concatenate(probs))
    df_prob_h=df_prob_h[0]
    
    # Away / Not Away
    df_a=df.replace({'H': 'NA', 'D': 'NA'})
    df_a=df_a[variables]
    if model=='XGBoost':
        le = LabelEncoder()
        df_a['FTR'] = le.fit_transform(df_a['FTR'])
    probs,preds=get_forecasts(df=df_a,features=features,stand_df=False,fct_num=8,fct_win=1,model=model)

    df_prob_a=pd.DataFrame(np.concatenate(probs))
    df_prob_a=df_prob_a[0]

    # Draw / Not Draw
    df_d=df.replace({'H': 'ND', 'A': 'ND'})
    df_d=df_d[variables]
    if model=='XGBoost':
        le = LabelEncoder()
        df_d['FTR'] = le.fit_transform(df_d['FTR'])
    probs,preds=get_forecasts(df=df_d,features=features,stand_df=False,fct_num=8,fct_win=1,model=model)

    df_prob_d=pd.DataFrame(np.concatenate(probs))
    df_prob_d=df_prob_d[0]


    df_3binomial=pd.DataFrame({'H': df_prob_h, 'A': df_prob_a,'D':df_prob_d})

    def find_winner(row):
        if ((row['H']>=row['A'])&(row['H']>=row['D'])):
            return 'H'
        elif ((row['A']>row['H'])&(row['A']>row['D'])):
            return 'A'
        else:
            return 'D'
            

    df_3binomial['final_pred']=df_3binomial.apply(lambda x: find_winner(x), axis=1)

    return df_3binomial



########################################################
# functions to generate some plots like tpr vs fpr, AUC curve 
# these functions i copied somewhere from Internet (don't remember the author)


def calculate_tpr_fpr(y_real, y_pred):
 
    cm = confusion_matrix(y_real, y_pred)
    TN = cm[0, 0]
    FP = cm[0, 1]
    FN = cm[1, 0]
    TP = cm[1, 1]
    
    tpr =  TP/(TP + FN) 
    fpr = 1 - TN/(TN+FP) 
    
    return tpr, fpr

def get_all_roc_coordinates(y_real, y_proba):

    tpr_list = [0]
    fpr_list = [0]
    for i in range(len(y_proba)):
        threshold = y_proba[i]
        y_pred = y_proba >= threshold
        tpr, fpr = calculate_tpr_fpr(y_real, y_pred)
        tpr_list.append(tpr)
        fpr_list.append(fpr)
    return tpr_list, fpr_list

def plot_roc_curve(tpr, fpr, scatter = True, ax = None):

    if ax == None:
        plt.figure(figsize = (5, 5))
        ax = plt.axes()
    
    if scatter:
        sns.scatterplot(x = fpr, y = tpr, ax = ax)
    sns.lineplot(x = fpr, y = tpr, ax = ax)
    sns.lineplot(x = [0, 1], y = [0, 1], color = 'green', ax = ax)
    plt.xlim(-0.05, 1.05)
    plt.ylim(-0.05, 1.05)
    plt.xlabel("FPR")
    plt.ylabel("TPR")

def multi_auprc(y_true_cat, y_score):
    y_true = OneHotEncoder().fit_transform(y_true_cat.reshape(-1, 1)).toarray()
    
    return average_precision_score(y_true, y_score)



######################################################
# convert grid to csv file with predictions
def grid_to_predictions(grid,title,y_test_adh,X_test):
    
    y_proba=grid.best_estimator_.predict_proba(X_test)
    y_pred=grid.best_estimator_.predict(X_test)

    df_predictions=pd.DataFrame(data=y_proba,columns=['A','D','H'])

    def find_winner_prob(row):
            if ((row['H']>=row['A'])&(row['H']>=row['D'])):
                return row['H']
            elif ((row['A']>row['H'])&(row['A']>row['D'])):
                return row['A']
            else:
                return row['D']

    def find_winner(row):
            if ((row['H']>=row['A'])&(row['H']>=row['D'])):
                return 'H'
            elif ((row['A']>row['H'])&(row['A']>row['D'])):
                return 'A'
            else:
                return 'D'       

    df_predictions['pred_prob']=df_predictions.apply(lambda x: find_winner_prob(x), axis=1)
    y_test_adh=y_test_adh.reset_index(drop=True)
    df_predictions['real_result']=y_test_adh  
    df_predictions['final_pred']=df_predictions.apply(lambda x: find_winner(x), axis=1)

    df_predictions.to_csv('predictions_'+title+'.csv')


# calculate RPS measure
def rps(y,y_pred):

    df_y = pd.DataFrame({'result': y})
    y = pd.get_dummies(df_y['result'])
    ncat = y_pred.shape[1] - 1
    return (
        np.mean(
            np.sum(
                (np.cumsum(y_pred, axis=1) - np.cumsum(y, axis=1)) ** 2, axis=1
            ) / ncat
        )
    )



##########################################################
# functions needed to generate results of every betting strategy I used

def find_winner_prob(row):
        if ((row['H']>=row['A'])&(row['H']>=row['D'])):
            return row['H']
        elif ((row['A']>row['H'])&(row['A']>row['D'])):
            return row['A']
        else:
            return row['D']

def find_winner(row):
        if ((row['H']>=row['A'])&(row['H']>=row['D'])):
            return 'H'
        elif ((row['A']>row['H'])&(row['A']>row['D'])):
            return 'A'
        else:
            return 'D'
        
def find_real_winner_odds(row):
        if (row['real_result']=='H'):
            return row['OddsH']
        elif (row['real_result']=='A'):
            return row['OddsA']
        else:
            return row['OddsD']

def find_pred_winner_odds(row):
        if (row['final_pred']=='H'):
            return row['OddsH']
        elif (row['final_pred']=='A'):
            return row['OddsA']
        else:
            return row['OddsD']

def find_profit(row,tax=1):
        if (row['final_pred']==row['real_result']):
            if row['stake']<0:
                return 0
            else:
                return ((row['winner_odds']*row['stake']*tax)-row['stake'])
        else:
            if row['stake']<0:
                return 0
            else:
                return -row['stake']
            
def find_profit_kelly(row):
        if (row['final_pred']==row['real_result']):
            if row['bankroll%']<0:
                return 0
            else:
                return (row['winner_odds']*row['stake'])-row['stake']
        else:
            if row['bankroll%']<0:
                return 0
            else:
                return -row['stake']

def find_profit_with_taxes_kelly(row):
    if (row['final_pred']==row['real_result']):
        if row['bankroll%']<0:
            return 0
        else:
            return (row['winner_odds']*row['stake_with_taxes']*0.88)-row['stake_with_taxes']
    else:
        if row['bankroll%']<0:
            return 0
        else:
            return -row['stake_with_taxes']

def find_profit_HNH(row,tax=False):
    if tax==False:
        if (row['H/NH_pred']=='NH'):
            if (row['real_result']!='H'):
                return (row['OddsX2']*row['stake'])-row['stake']
            else:
                return -row['stake']
        else:
            return 0
    elif tax==True:
        if (row['H/NH_pred']=='NH'):
            if (row['real_result']!='H'):
                return (row['OddsX2']*row['stake']*0.88)-row['stake']
            else:
                return -row['stake']
        else:
            return 0

def find_profit_ANA(row,tax=False):
    if tax==False:
        if (row['A/NA_pred']=='NA'):
            if (row['real_result']!='A'):
                return (row['Odds1X']*row['stake'])-row['stake']
            else:
                return -row['stake']
        else:
            return 0
    elif tax==True:
        if (row['A/NA_pred']=='NA'):
            if (row['real_result']!='A'):
                return (row['Odds1X']*row['stake']*0.88)-row['stake']
            else:
                return -row['stake']
        else:
            return 0
            
def find_profit_DND(row,tax=False):
    if tax==False:
        if (row['D/ND_pred']=='ND'):
            if (row['real_result']!='D'):
                return (row['Odds12']*row['stake'])-row['stake']
            else:
                return -row['stake']
        else:
            return 0
    elif tax==True:
        if (row['D/ND_pred']=='ND'):
            if (row['real_result']!='D'):
                return (row['Odds12']*row['stake']*0.88)-row['stake']
            else:
                return -row['stake']
        else:
            return 0